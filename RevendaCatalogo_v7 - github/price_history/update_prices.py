#!/usr/bin/env python3
"""
update_prices.py — histórico de preços a partir do CSV de um fornecedor.

Fluxo:
  1. Localiza o CSV (argumento --csv, CSV_SOURCE, ou o upload mais recente da app catalogo)
  2. Lê e valida o CSV
  3. Compara com a versão anterior (output/prices_current.csv)
  4. Faz backup da versão anterior em output/backups/
  5. Gera Excel formatado (Produtos + Análise + Mudanças)
  6. Regista em output/update_log.txt

Uso:
  python update_prices.py --supplier <chave>   # upload mais recente desse fornecedor na app
  python update_prices.py --csv caminho.csv
  CSV_SOURCE=... OUTPUT_DIR=... python update_prices.py
"""
from __future__ import annotations

import argparse
import glob
import os
import shutil
import statistics
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Optional

import pandas as pd

# Consola do Windows usa cp1252 por omissão e rebentava com "✓"; UTF-8 com substituição
# garante que nem o terminal nem o Task Scheduler falham por causa de um símbolo.
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8", errors="replace")

# ── Configuração ajustável ────────────────────────────────────────────────────
BASE_DIR = Path(__file__).resolve().parent
# Caminho do CSV. Vazio ⇒ procura o upload mais recente da app catalogo.
CSV_SOURCE: str = os.getenv("CSV_SOURCE", "")
# Chave do fornecedor na app (prefixo dos ficheiros em uploads/). Vazio ⇒ qualquer CSV.
SUPPLIER: str = os.getenv("SUPPLIER", "")
# Pasta onde ficam o CSV/Excel atuais, o log e os backups.
OUTPUT_DIR: Path = Path(os.getenv("OUTPUT_DIR", BASE_DIR / "output"))
# Pasta de uploads da app (usada só quando CSV_SOURCE está vazio).
APP_UPLOADS: Path = BASE_DIR.parent / "catalogo" / "uploads"

CURRENT_CSV_NAME = "prices_current.csv"
CURRENT_XLSX_NAME = "prices_current.xlsx"
LOG_NAME = "update_log.txt"

REQUIRED_COLS = ["id", "product_name", "price"]
# Colunas opcionais: se o CSV não as tiver, são criadas vazias / a zero.
OPTIONAL_TEXT_COLS = ["ean", "brand", "product_url", "product_type", "sku"]
OPTIONAL_NUM_COLS = ["supply_status", "stock_amount"]
KEY_COL = "id"          # tem de ser único e estável entre versões (o EAN pode ter vazios)


# ── Resultado da deteção de mudanças ──────────────────────────────────────────
@dataclass
class Changes:
    """Diferenças entre a versão anterior e a atual do CSV."""
    novos: pd.DataFrame = field(default_factory=pd.DataFrame)
    removidos: pd.DataFrame = field(default_factory=pd.DataFrame)
    precos: pd.DataFrame = field(default_factory=pd.DataFrame)   # id, product_name, preco_anterior, preco_atual, variacao_pct
    tem_anterior: bool = False

    @property
    def total(self) -> int:
        return len(self.novos) + len(self.removidos) + len(self.precos)


# ── Localização e leitura ─────────────────────────────────────────────────────
def find_csv(explicit: str = "", supplier: str = SUPPLIER) -> Path:
    """Devolve o CSV a processar. Ordem: argumento → CSV_SOURCE → upload mais recente da app."""
    for cand in (explicit, CSV_SOURCE):
        if cand:
            p = Path(cand).expanduser()
            if not p.is_file():
                raise FileNotFoundError(f"CSV não encontrado: {p}")
            return p
    pattern = f"{supplier}_*.csv" if supplier else "*.csv"
    uploads = sorted(glob.glob(str(APP_UPLOADS / pattern)), key=os.path.getmtime, reverse=True)
    if not uploads:
        raise FileNotFoundError(
            f"Nenhum CSV indicado e nenhum upload '{pattern}' encontrado em "
            f"{APP_UPLOADS}. Usa --supplier <chave>, --csv caminho.csv ou CSV_SOURCE")
    return Path(uploads[0])


def load_csv(path: Path) -> pd.DataFrame:
    """Lê o CSV com tolerância a encoding e valida as colunas."""
    last_err: Optional[Exception] = None
    for enc in ("utf-8", "utf-8-sig", "cp1252", "latin1"):
        try:
            df = pd.read_csv(path, encoding=enc, dtype={"ean": str, "sku": str})
            break
        except (UnicodeDecodeError, ValueError) as exc:
            last_err = exc
    else:
        raise ValueError(f"Não foi possível ler {path}: {last_err}")

    df.columns = [c.strip().lower() for c in df.columns]
    missing = [c for c in REQUIRED_COLS if c not in df.columns]
    if missing:
        raise ValueError(f"CSV sem as colunas esperadas: {missing}. Colunas presentes: {list(df.columns)}")
    for col in OPTIONAL_TEXT_COLS:
        if col not in df.columns:
            df[col] = ""
    for col in OPTIONAL_NUM_COLS:
        if col not in df.columns:
            df[col] = 0

    df["price"] = pd.to_numeric(df["price"], errors="coerce")
    df["stock_amount"] = pd.to_numeric(df["stock_amount"], errors="coerce").fillna(0).astype(int)
    df["supply_status"] = pd.to_numeric(df["supply_status"], errors="coerce").fillna(0).astype(int)
    df[KEY_COL] = df[KEY_COL].astype(str).str.strip()
    df = df.dropna(subset=["price"]).drop_duplicates(subset=[KEY_COL], keep="first")
    return df.reset_index(drop=True)


def load_previous(output_dir: Path) -> Optional[pd.DataFrame]:
    """Versão anterior (se existir). Não existir é normal na 1.ª execução."""
    prev = output_dir / CURRENT_CSV_NAME
    if not prev.is_file():
        return None
    try:
        return load_csv(prev)
    except Exception as exc:  # anterior corrompido não deve travar a execução
        print(f"  ! Versão anterior ilegível ({exc}); a continuar sem comparação")
        return None


# ── Deteção de mudanças ───────────────────────────────────────────────────────
def detect_changes(prev: Optional[pd.DataFrame], cur: pd.DataFrame) -> Changes:
    """Novos, removidos e mudanças de preço entre duas versões (chave: id)."""
    if prev is None:
        return Changes()
    p = prev.set_index(KEY_COL)
    c = cur.set_index(KEY_COL)
    novos = cur[~cur[KEY_COL].isin(p.index)][["id", "product_name", "brand", "price"]]
    removidos = prev[~prev[KEY_COL].isin(c.index)][["id", "product_name", "brand", "price"]]
    comum = p.index.intersection(c.index)
    j = pd.DataFrame({
        "product_name": c.loc[comum, "product_name"],
        "preco_anterior": p.loc[comum, "price"],
        "preco_atual": c.loc[comum, "price"],
    })
    j = j[j["preco_anterior"] != j["preco_atual"]].copy()
    j["variacao_pct"] = ((j["preco_atual"] - j["preco_anterior"]) / j["preco_anterior"] * 100).round(2)
    j = j.reset_index().rename(columns={"index": "id"})
    return Changes(novos=novos.reset_index(drop=True), removidos=removidos.reset_index(drop=True),
                   precos=j.sort_values("variacao_pct", ascending=False).reset_index(drop=True),
                   tem_anterior=True)


# ── Backup e log ──────────────────────────────────────────────────────────────
def backup_previous(output_dir: Path, ts: str) -> Optional[Path]:
    """Copia o CSV atual para backups/ antes de o substituir."""
    prev = output_dir / CURRENT_CSV_NAME
    if not prev.is_file():
        return None
    bdir = output_dir / "backups"
    bdir.mkdir(parents=True, exist_ok=True)
    dest = bdir / f"backup_{ts}.csv"
    shutil.copy2(prev, dest)
    return dest


def write_log(output_dir: Path, lines: list[str]) -> None:
    with open(output_dir / LOG_NAME, "a", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")


# ── Excel ─────────────────────────────────────────────────────────────────────
EURO_FMT = '"€"#,##0.00'


def _autowidth(ws) -> None:
    from openpyxl.utils import get_column_letter
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 8), 70)


def _style_header(ws) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _write_df(wb, title: str, df: pd.DataFrame, euro_cols: tuple[str, ...] = ()) -> None:
    from openpyxl.utils.dataframe import dataframe_to_rows
    ws = wb.create_sheet(title)
    for row in dataframe_to_rows(df, index=False, header=True):
        ws.append(row)
    if df.empty:
        return
    _style_header(ws)
    for name in euro_cols:
        if name in df.columns:
            idx = list(df.columns).index(name) + 1
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=idx).number_format = EURO_FMT
    _autowidth(ws)


def build_analysis(df: pd.DataFrame) -> list[tuple]:
    """Linhas (rótulo, valor) para a aba Análise."""
    precos = df["price"].dropna()
    em_stock = int((df["stock_amount"] > 0).sum())
    rows: list[tuple] = [
        ("Total de produtos", len(df)),
        ("Produtos em stock", em_stock),
        ("Produtos indisponíveis", len(df) - em_stock),
        ("Preço mínimo", float(precos.min()) if len(precos) else 0),
        ("Preço médio", round(float(precos.mean()), 2) if len(precos) else 0),
        ("Preço mediana", float(statistics.median(precos)) if len(precos) else 0),
        ("Preço máximo", float(precos.max()) if len(precos) else 0),
        ("", ""),
        ("TOP 15 CATEGORIAS", "Produtos"),
    ]
    rows += [(k, int(v)) for k, v in df["product_type"].fillna("(sem tipo)").value_counts().head(15).items()]
    rows += [("", ""), ("TOP 10 MARCAS", "Produtos")]
    rows += [(k, int(v)) for k, v in df["brand"].fillna("(sem marca)").value_counts().head(10).items()]
    return rows


def write_excel(df: pd.DataFrame, changes: Changes, dest: Path, ts_human: str) -> None:
    """Excel com abas Produtos (ordenado por categoria), Análise e Mudanças."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)

    produtos = df.sort_values(["product_type", "brand", "product_name"], na_position="last")
    produtos = produtos[["product_type", "brand", "product_name", "price", "stock_amount",
                         "supply_status", "ean", "sku", "id", "product_url"]]
    produtos = produtos.rename(columns={
        "product_type": "Categoria", "brand": "Marca", "product_name": "Produto",
        "price": "Preço (€)", "stock_amount": "Stock", "supply_status": "Disponível",
        "ean": "EAN", "sku": "SKU", "id": "ID", "product_url": "URL"})
    _write_df(wb, "Produtos", produtos, euro_cols=("Preço (€)",))

    ws = wb.create_sheet("Análise")
    ws.append([f"Análise — {ts_human}", ""])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Indicador", "Valor"])
    for label, value in build_analysis(df):
        ws.append([label, value])
        if label.startswith("Preço") and isinstance(value, (int, float)):
            ws.cell(row=ws.max_row, column=2).number_format = EURO_FMT
        if label.startswith("TOP"):
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for cell in ws[2]:
        cell.font = Font(bold=True)
    _autowidth(ws)

    if changes.tem_anterior:
        resumo = pd.DataFrame({
            "Indicador": ["Produtos novos", "Produtos removidos", "Mudanças de preço",
                          "Subidas de preço", "Descidas de preço"],
            "Valor": [len(changes.novos), len(changes.removidos), len(changes.precos),
                      int((changes.precos["variacao_pct"] > 0).sum()) if len(changes.precos) else 0,
                      int((changes.precos["variacao_pct"] < 0).sum()) if len(changes.precos) else 0],
        })
        _write_df(wb, "Mudanças", resumo)
        if len(changes.precos):
            pc = changes.precos.rename(columns={"product_name": "Produto", "preco_anterior": "Preço anterior (€)",
                                                "preco_atual": "Preço atual (€)", "variacao_pct": "Variação %"})
            _write_df(wb, "Preços alterados", pc, euro_cols=("Preço anterior (€)", "Preço atual (€)"))
        if len(changes.novos):
            _write_df(wb, "Novos", changes.novos.rename(columns={"product_name": "Produto", "price": "Preço (€)"}),
                      euro_cols=("Preço (€)",))
        if len(changes.removidos):
            _write_df(wb, "Removidos", changes.removidos.rename(columns={"product_name": "Produto", "price": "Preço (€)"}),
                      euro_cols=("Preço (€)",))

    wb.save(dest)


# ── Orquestração ──────────────────────────────────────────────────────────────
def run(csv_arg: str = "", output_dir: Path = OUTPUT_DIR,
        supplier: str = SUPPLIER) -> tuple[pd.DataFrame, Changes]:
    """Executa o pipeline completo; devolve (dados atuais, mudanças). Levanta exceção só se o CSV falhar."""
    now = datetime.now()
    ts, ts_human = now.strftime("%Y%m%d_%H%M%S"), now.strftime("%d/%m/%Y %H:%M:%S")
    output_dir.mkdir(parents=True, exist_ok=True)
    log: list[str] = [f"[{ts_human}] ---- atualização ----"]

    src = find_csv(csv_arg, supplier)
    df = load_csv(src)
    print(f"✓ CSV processado com {len(df)} produtos ({src.name})")
    log.append(f"  CSV: {src} ({len(df)} produtos)")

    prev = load_previous(output_dir)
    changes = detect_changes(prev, df)
    if changes.tem_anterior:
        print(f"✓ Mudanças detectadas: {changes.total} "
              f"(novos {len(changes.novos)}, removidos {len(changes.removidos)}, preços {len(changes.precos)})")
        log.append(f"  Mudanças: novos={len(changes.novos)} removidos={len(changes.removidos)} preços={len(changes.precos)}")
        pc = changes.precos
        for etiqueta, sub in (("Top 5 aumentos", pc[pc["variacao_pct"] > 0].head(5)),
                              ("Top 5 reduções", pc[pc["variacao_pct"] < 0].sort_values("variacao_pct").head(5))):
            if len(sub):
                log.append(f"  {etiqueta}:")
                log += [f"    {r.product_name[:60]:60} {r.preco_anterior:9.2f} → {r.preco_atual:9.2f} ({r.variacao_pct:+.1f}%)"
                        for r in sub.itertuples()]
    else:
        print("• Sem versão anterior — comparação ignorada (normal na 1.ª execução)")
        log.append("  Sem versão anterior")

    bk = backup_previous(output_dir, ts)
    if bk:
        print(f"✓ Backup criado: {bk.name}")
        log.append(f"  Backup: {bk.name}")

    df.to_csv(output_dir / CURRENT_CSV_NAME, index=False, encoding="utf-8")

    xlsx = output_dir / CURRENT_XLSX_NAME
    try:
        write_excel(df, changes, xlsx, ts_human)
        print(f"✓ Excel criado: {xlsx.name}")
        log.append(f"  Excel: {xlsx.name}")
    except Exception as exc:   # Excel é secundário: avisar e continuar
        print(f"  ! Excel falhou ({exc}) — CSV atualizado na mesma")
        log.append(f"  Excel FALHOU: {exc}")

    write_log(output_dir, log)
    print(f"✓ Log atualizado: {LOG_NAME}")
    return df, changes


def main(argv: Optional[list[str]] = None) -> int:
    ap = argparse.ArgumentParser(description="Atualiza o Excel de histórico de preços a partir de um CSV.")
    ap.add_argument("--csv", default="", help="caminho do CSV (senão CSV_SOURCE ou upload mais recente da app)")
    ap.add_argument("--supplier", default=SUPPLIER,
                    help="chave do fornecedor na app catalogo (usa o upload mais recente desse fornecedor)")
    ap.add_argument("--out", default="", help=f"pasta de output (default: {OUTPUT_DIR})")
    args = ap.parse_args(argv)
    try:
        run(args.csv, Path(args.out) if args.out else OUTPUT_DIR, args.supplier)
        return 0
    except FileNotFoundError as exc:
        print(f"ERRO: {exc}", file=sys.stderr)
        return 1
    except ValueError as exc:
        print(f"ERRO no CSV: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
